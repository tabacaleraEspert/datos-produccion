"""
Parser de ConsumoTeoricoTabaco y ConsumoRealTabaco desde la Minuta Excel.

Tabla destino: RendimientoTabaco (fecha, ConsumoTeoricoKg, ConsumoRealKg)

Las celdas varían por versión de minuta. Se usan alias en config.RENDIMIENTO_TABACO_CELLS:
- tabaco_consumo_teorico → ConsumoTeoricoKg
- tabaco_consumo_real → ConsumoRealKg

Con use_excel_names=True se leen los nombres definidos en Excel (Name Manager).
Si agregan filas arriba, Excel actualiza la ref y el script sigue funcionando.

fecha = lunes de la semana anterior a la fecha de la minuta.
"""
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

from ..parse_minuta import MinutaFileInfo
from config import RENDIMIENTO_TABACO_CELLS

SHEET = "Datos crudos"

ALIAS_TEORICO = "tabaco_consumo_teorico"
ALIAS_REAL = "tabaco_consumo_real"


def _parse_cell_ref(cell: str) -> Optional[Tuple[int, int]]:
    """Convierte celda Excel (ej. B83, H83) en (fila_0based, col_0based)."""
    m = re.match(r"^([A-Za-z]+)(\d+)$", str(cell).strip())
    if not m:
        return None
    col_str, row_str = m.groups()
    col = 0
    for c in col_str.upper():
        col = col * 26 + (ord(c) - ord("A") + 1)
    col -= 1  # 0-based
    row = int(row_str) - 1  # 1-based -> 0-based
    return (row, col)


@dataclass
class RendimientoTabacoRow:
    """Una fila para RendimientoTabaco."""
    fecha: str  # YYYY-MM-DD (lunes de la semana)
    consumo_teorico_kg: float
    consumo_real_kg: float


@dataclass
class RendimientoTabacoParseResult:
    """Resultado del parse."""
    fila: Optional[RendimientoTabacoRow] = None
    errors: List[str] = field(default_factory=list)


def _parse_numero(cell) -> Optional[float]:
    """Convierte celda a float. None si vacío o no numérico."""
    if cell is None or (hasattr(pd, "isna") and pd.isna(cell)):
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    try:
        return float(str(cell).strip().replace(",", "."))
    except (ValueError, TypeError):
        return None


def _read_excel_named_value(xlsx_path: str, name: str) -> Optional[float]:
    """Lee el valor de un nombre definido en Excel. None si no existe o es #REF!."""
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        if name not in wb.defined_names:
            return None
        defn = wb.defined_names[name]
        for sheet_title, coord in defn.destinations:
            ws = wb[sheet_title]
            # Si es rango (ej. B82:B82), tomar primera celda
            cell_ref = coord.split(":")[0].replace("$", "")
            val = ws[cell_ref].value
            wb.close()
            return _parse_numero(val) if val is not None else None
        wb.close()
    except Exception:
        pass
    return None


def parse_rendimiento_tabaco(
    xlsx_path: str,
    minuta_info: MinutaFileInfo,
    sheet_name: Optional[str] = None,
) -> RendimientoTabacoParseResult:
    """Lee tabaco_consumo_teorico y tabaco_consumo_real según la versión de la minuta."""
    result = RendimientoTabacoParseResult()
    if not minuta_info.fecha:
        result.errors.append("Falta fecha en minuta_info.")
        return result

    version_key = (minuta_info.version or "").lower()
    cells = RENDIMIENTO_TABACO_CELLS.get(version_key)
    if not cells:
        result.errors.append(
            f"Versión '{minuta_info.version}' sin celdas definidas en RENDIMIENTO_TABACO_CELLS. "
            f"Versiones disponibles: {list(RENDIMIENTO_TABACO_CELLS.keys())}"
        )
        return result

    cell_teorico = cells.get(ALIAS_TEORICO)
    cell_real = cells.get(ALIAS_REAL)
    sheet_override = cells.get("sheet")
    use_excel_names = cells.get("use_excel_names", False)
    if not cell_teorico or not cell_real:
        result.errors.append(
            f"Faltan alias {ALIAS_TEORICO} o {ALIAS_REAL} para versión {version_key}."
        )
        return result

    teorico: Optional[float] = None
    real: Optional[float] = None
    source_info = ""

    if use_excel_names:
        teorico = _read_excel_named_value(xlsx_path, ALIAS_TEORICO)
        real = _read_excel_named_value(xlsx_path, ALIAS_REAL)
        if teorico is not None or real is not None:
            source_info = "nombres Excel"

    if teorico is None and real is None:
        coords_teorico = _parse_cell_ref(cell_teorico)
        coords_real = _parse_cell_ref(cell_real)
        if not coords_teorico or not coords_real:
            result.errors.append(
                f"Celdas inválidas: {cell_teorico}, {cell_real}. Usar formato Excel (ej. B83, H83)."
            )
            return result

        sheet = sheet_name or sheet_override or SHEET
        try:
            df = pd.read_excel(xlsx_path, sheet_name=sheet, header=None, engine="openpyxl")
        except Exception as e:
            result.errors.append(f"No se pudo leer la hoja '{sheet}': {e}")
            return result

        row_idx_t, col_idx_t = coords_teorico
        row_idx_r, col_idx_r = coords_real

        if row_idx_t < len(df):
            r = df.iloc[row_idx_t]
            teorico = _parse_numero(r.iloc[col_idx_t] if col_idx_t < len(r) else None)
        if row_idx_r < len(df):
            r = df.iloc[row_idx_r]
            real = _parse_numero(r.iloc[col_idx_r] if col_idx_r < len(r) else None)
        source_info = f"celdas {cell_teorico}, {cell_real}"

    if teorico is None:
        result.errors.append(
            f"{ALIAS_TEORICO} vacío o no numérico ({source_info or 'sin datos'})."
        )
    if real is None:
        result.errors.append(f"{ALIAS_REAL} vacío o no numérico ({source_info or 'sin datos'}).")

    if teorico is None and real is None:
        return result

    # Lunes de la semana anterior a la fecha de la minuta
    d = datetime.strptime(minuta_info.fecha, "%Y-%m-%d").date()
    lunes_actual = d - timedelta(days=d.weekday())
    lunes_anterior = lunes_actual - timedelta(days=7)
    fecha = lunes_anterior.strftime("%Y-%m-%d")

    result.fila = RendimientoTabacoRow(
        fecha=fecha,
        consumo_teorico_kg=teorico if teorico is not None else 0.0,
        consumo_real_kg=real if real is not None else 0.0,
    )
    return result
