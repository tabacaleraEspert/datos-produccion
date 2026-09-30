"""Layout UNIFICADO (ago-2026+): un solo Excel "Plan NNN ... (Villafranca-Ale)"
con hojas `Minuta`, `Datos crudos`, `Plan - Semanal` y `REND.TABACO`.

Reemplaza al par Plan(v22)+Minuta(v12) para las semanas nuevas. Mapea a las
mismas 5 tablas de Azure. Fecha de las tablas semanales = lunes de la semana.

Estructura relevada (Plan 290):
- Plan - Semanal: fechas en fila 112 (C..J); bloques de 13 filas por marca
  desde f113: +0..+3 = plan (variantes), +4..+7 = real, +8 = "Total Real"
  (se usa para validar), +9..+12 diferencias.
- Minuta: eficiencias en filas ~42-49, label "MK9-1 [67% @ 4300]",
  B/C = TD + valor (fracción 0-1), D/E = TJ + valor. Tabaco: fila con
  A='Teórico (Kg)' → B teórico, H real.
- Datos crudos: máquinas por fila (label col B con la unidad), días en
  tríos de columnas desde C: (TD,TJ,TN) × Lun..Dom. 'N'/'-' = sin dato.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

# --- catálogos (mismos ids que el resto del ETL / DB) -----------------------

# hoja Minuta: label de eficiencia → Id_Maquina
EFIC_MAQ = {
    "mk9-1": 1, "mk9 - 1": 1,
    "mk9-2": 2, "mk9 - 2": 2,
    "protos": 13,
    "hlp 1": 3, "hlp1": 3,
    "hlp 2": 4, "hlp2": 4,
    "amf 3000 - paper": 6,
    "paper -": 15, "paper ": 15,
    "gd1 - amf": 16,
    "gd2 - amf": 17,
}

# hoja Datos crudos: (patrón de label col B, Id_Maquina). El patrón matchea el
# label normalizado (minúsculas, sin saltos de línea). Orden importa: el primer
# patrón que matchea gana y cada id se toma una sola vez.
CRUDOS_MAQ: List[Tuple[str, int]] = [
    ("mk9 - 1", 1),
    ("mk9 - 2", 2),
    ("protos", 13),
    ("hlp 1", 3),
    ("hlp 2", 4),
    ("amf 3000 - gd1", 7),
    ("amf 3000 - gd2", 5),
    ("amf 5000 - gd1", 8),
    ("amf 5000 - gd2", 9),
    ("paper - amf 3000", 15),
    ("amf 3000", 6),           # después de los AMF 3000 específicos
    ("gd1-5000/3000", 16),
    ("gd2-5000/3000", 17),
    ("filtrera polaris (varas)", 10),
    ("filtrera molins 1c (varas)", 14),
    ("filtrera molins 2c (varas)", 18),
]
# prefijos de labels que NO son la fila principal de producción de una máquina
CRUDOS_EXCLUIR_PREFIJO = ("batea", "cajas", "rechazo", "descarte", "total")

TURNOS = ("TD", "TJ", "TN")


def _norm(s: Any) -> str:
    if s is None:
        return ""
    s = str(s).replace("\n", " ").strip().lower()
    return re.sub(r"\s+", " ", s)


def _num(v: Any) -> Optional[float]:
    """Número o None ('N', '-', '', texto = sin dato)."""
    if v is None or isinstance(v, str):
        try:
            return float(str(v).replace(",", ".")) if v not in (None, "", "N", "-", "?") else None
        except (TypeError, ValueError):
            return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


# --- Plan - Semanal ----------------------------------------------------------

def _fechas_semana(ws) -> List[Tuple[int, date]]:
    """(columna, fecha) de la fila 112. La primera es el lunes."""
    out = []
    for col in range(3, 12):  # C..K
        v = ws.cell(row=112, column=col).value
        if hasattr(v, "date"):
            out.append((col, v.date()))
        elif isinstance(v, date):
            out.append((col, v))
    return out


def _marca_base(label: str) -> str:
    """'Melbourne Red TABES FLAMA 6.5' → 'Melbourne Red' (corta en TABES/Total)."""
    s = re.split(r"\s+TABES\b", label, flags=re.I)[0]
    s = re.sub(r"\s+(Total Real|Real Total)\s*$", "", s, flags=re.I)
    return s.strip()


def parse_plan(ws, buscar_id_marca) -> Tuple[List[tuple], List[tuple], List[str]]:
    """→ (plan_rows, real_rows, avisos) con rows = (fecha, id_marca, cantidad)."""
    fechas = _fechas_semana(ws)
    if not fechas:
        raise ValueError("Plan - Semanal: no encontré fechas en la fila 112")
    plan: Dict[Tuple[date, int], int] = {}
    real: Dict[Tuple[date, int], int] = {}
    avisos: List[str] = []

    f = 113
    while f < ws.max_row:
        label = ws.cell(row=f, column=2).value
        if not label or not str(label).strip():
            break
        base = _marca_base(str(label))
        if _norm(base).startswith(("sub total", "total", "diferencia")):
            break
        id_marca = buscar_id_marca(base)
        if id_marca is None:
            avisos.append(f"marca sin match: '{base}' (fila {f})")
        else:
            total_real_chk = 0.0
            for col, fecha in fechas:
                p = sum(_num(ws.cell(row=f + o, column=col).value) or 0 for o in range(0, 4))
                r = sum(_num(ws.cell(row=f + o, column=col).value) or 0 for o in range(4, 8))
                if p:
                    plan[(fecha, id_marca)] = plan.get((fecha, id_marca), 0) + int(round(p))
                if r:
                    real[(fecha, id_marca)] = real.get((fecha, id_marca), 0) + int(round(r))
                total_real_chk += r
            tot_fila = sum(_num(ws.cell(row=f + 8, column=col).value) or 0 for col, _ in fechas)
            if tot_fila and abs(tot_fila - total_real_chk) > 0.5:
                avisos.append(
                    f"'{base}': real sumado {total_real_chk:.0f} ≠ fila Total {tot_fila:.0f}"
                )
        f += 13

    p_rows = [(fe, m, c) for (fe, m), c in sorted(plan.items())]
    r_rows = [(fe, m, c) for (fe, m), c in sorted(real.items())]
    return p_rows, r_rows, avisos


# --- Minuta: eficiencias + tabaco -------------------------------------------

_RE_EFIC = re.compile(r"^(.*?)\s*\[(\d+(?:\.\d+)?)%\s*@\s*(\d+)\]")


def parse_minuta(ws, lunes: date) -> Tuple[List[tuple], Optional[tuple], List[str]]:
    """→ (rend_rows, tabaco_row, avisos).

    rend_rows = (fecha, id_maquina, turno, eficiencia_pct, rend_deseado, vel_deseada)
    tabaco_row = (fecha, teorico_kg, real_kg) o None
    """
    rend: List[tuple] = []
    avisos: List[str] = []
    tabaco = None
    for f in range(1, min(ws.max_row, 120) + 1):
        a = ws.cell(row=f, column=1).value
        if not a:
            continue
        m = _RE_EFIC.match(str(a).strip())
        if m:
            nombre, obj, vel = _norm(m.group(1)), float(m.group(2)), int(m.group(3))
            mid = next((i for pat, i in EFIC_MAQ.items() if nombre.startswith(pat)), None)
            if mid is None:
                avisos.append(f"eficiencia sin máquina: '{m.group(1)}' (fila {f})")
                continue
            for col_t, col_v in ((2, 3), (4, 5)):  # (B,C)=TD  (D,E)=TJ
                t = _norm(ws.cell(row=f, column=col_t).value).upper().replace(":", "")
                v = _num(ws.cell(row=f, column=col_v).value)
                if t in ("TD", "TJ", "TN") and v is not None:
                    pct = v * 100 if v <= 1.5 else v
                    rend.append((lunes, mid, t, round(pct, 1), obj, vel))
        norm_a = _sin_acentos(_norm(a))
        if norm_a.startswith("teorico (kg)"):
            teo = _num(ws.cell(row=f, column=2).value)
            rea = _num(ws.cell(row=f, column=8).value)
            if teo or rea:
                tabaco = (lunes, teo, rea)
    return rend, tabaco, avisos


# --- Datos crudos: producción por máquina/día/turno --------------------------

def parse_crudos(ws, lunes: date) -> Tuple[List[tuple], List[str]]:
    """→ (prod_rows, avisos) con rows = (fecha, id_maquina, turno, cantidad)."""
    usados: set[int] = set()
    filas: List[Tuple[int, int]] = []  # (fila, id)
    for f in range(1, min(ws.max_row, 120) + 1):
        label = _norm(ws.cell(row=f, column=2).value)
        if not label:
            continue
        if label.startswith(CRUDOS_EXCLUIR_PREFIJO):
            continue
        for pat, mid in CRUDOS_MAQ:
            if mid in usados:
                continue
            if pat in label:
                filas.append((f, mid))
                usados.add(mid)
                break
    avisos = []
    esperadas = {mid for _, mid in CRUDOS_MAQ}
    if esperadas - usados:
        avisos.append(f"datos crudos: máquinas no encontradas ids={sorted(esperadas - usados)}")

    prod: List[tuple] = []
    for f, mid in filas:
        for d in range(7):  # Lun..Dom
            fecha = lunes + timedelta(days=d)
            for t, off in zip(TURNOS, range(3)):
                v = _num(ws.cell(row=f, column=3 + d * 3 + off).value)
                if v is not None and v > 0:
                    prod.append((fecha, mid, t, int(round(v))))
    return prod, avisos


# --- workbook completo -------------------------------------------------------

def parse_unificado(path: str | Path, buscar_id_marca) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    faltan = {"Minuta", "Datos crudos", "Plan - Semanal"} - set(wb.sheetnames)
    if faltan:
        raise ValueError(f"{Path(path).name}: no es formato unificado (faltan hojas {faltan})")
    ws_plan = wb["Plan - Semanal"]
    fechas = _fechas_semana(ws_plan)
    if not fechas:
        raise ValueError(f"{Path(path).name}: sin fechas en Plan - Semanal fila 112")
    # Lunes CALENDARIO de la semana (los planes de semana corta arrancan martes,
    # pero la grilla de Datos crudos siempre es Lun..Dom).
    lunes = fechas[0][1] - timedelta(days=fechas[0][1].weekday())

    plan_rows, real_rows, av1 = parse_plan(ws_plan, buscar_id_marca)
    rend_rows, tabaco, av2 = parse_minuta(wb["Minuta"], lunes)
    prod_rows, av3 = parse_crudos(wb["Datos crudos"], lunes)
    wb.close()
    return {
        "archivo": Path(path).name,
        "lunes": lunes,
        "fechas": [fe for _, fe in fechas],
        "plan": plan_rows,
        "real": real_rows,
        "rendimiento": rend_rows,
        "tabaco": tabaco,
        "produccion_turno": prod_rows,
        "avisos": av1 + av2 + av3,
    }
